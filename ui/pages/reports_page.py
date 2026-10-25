from PySide6.QtCore import Qt
from PySide6.QtWidgets import (
    QAbstractItemView,
    QComboBox,
    QDialog,
    QDialogButtonBox,
    QFileDialog,
    QFrame,
    QGridLayout,
    QHBoxLayout,
    QHeaderView,
    QLabel,
    QLineEdit,
    QMessageBox,
    QPushButton,
    QTableWidget,
    QTableWidgetItem,
    QVBoxLayout,
    QWidget,
)
from sqlalchemy.exc import SQLAlchemyError

from database.db import init_db
from database.models.cafe.order_archive import OrderArchive

# ==========================================================
# پالت رنگی و استایل‌های مشترک
# ==========================================================
GREEN = "#28745A"
GREEN_DARK = "#1E5844"
GREEN_LIGHT = "#E8F3EB"
MUTED = "#6F7D76"
BORDER = "#E2E8E4"
WHITE = "#FFFFFF"
BG = "#F7FAF8"
TEXT = "#203028"
DANGER = "#D64545"
DANGER_LIGHT = "#FDE8E8"

PERSIAN_DIGITS = str.maketrans("0123456789", "۰۱۲۳۴۵۶۷۸۹")


def to_persian_digits(value) -> str:
    """تبدیل رقم‌ها به فارسی."""
    return str(value).translate(PERSIAN_DIGITS)


def money(value) -> str:
    """فرمت مبلغ به ریال/تومان با جداکننده هزارگان."""
    try:
        val = int(round(float(value or 0)))
    except (ValueError, TypeError):
        val = 0
    formatted = f"{val:,}".replace(",", "٬")
    return f"{to_persian_digits(formatted)} تومان"


def format_number(value) -> str:
    """فرمت عدد صحیح با ارقام فارسی."""
    try:
        val = int(value or 0)
    except (ValueError, TypeError):
        val = 0
    return to_persian_digits(f"{val:,}".replace(",", "٬"))


def customer_text(customer_name: str | None) -> str:
    return customer_name.strip() if customer_name and customer_name.strip() else "مشتری آزاد (مهمان)"


# ==========================================================
# ویجت وضعیت پرداخت (Badge)
# ==========================================================
class PaymentBadge(QWidget):
    """بج کپسولی وضعیت پرداخت درون سلول جدول."""

    def __init__(self, is_paid: bool, parent=None):
        super().__init__(parent)
        layout = QHBoxLayout(self)
        layout.setContentsMargins(6, 4, 6, 4)
        layout.setAlignment(Qt.AlignmentFlag.AlignCenter)

        label = QLabel("پرداخت‌شده" if is_paid else "پرداخت‌نشده")
        bg_color = GREEN_LIGHT if is_paid else DANGER_LIGHT
        fg_color = GREEN if is_paid else DANGER

        label.setStyleSheet(
            f"""
            QLabel {{
                background-color: {bg_color};
                color: {fg_color};
                border-radius: 9px;
                padding: 3px 10px;
                font-size: 11px;
                font-weight: 700;
            }}
            """
        )
        layout.addWidget(label)


# ==========================================================
# کارت‌های آماری بالای صفحه
# ==========================================================
class ReportStatCard(QFrame):
    def __init__(self, title: str, value: str = "۰", icon: str = "◈", accent: str = GREEN, parent=None):
        super().__init__(parent)
        self.setObjectName("statCard")
        self.setMinimumHeight(96)

        layout = QHBoxLayout(self)
        layout.setContentsMargins(16, 14, 16, 14)
        layout.setSpacing(14)

        icon_box = QLabel(icon)
        icon_box.setAlignment(Qt.AlignmentFlag.AlignCenter)
        icon_box.setFixedSize(44, 44)
        icon_box.setStyleSheet(
            f"""
            background: {accent}18;
            color: {accent};
            border-radius: 22px;
            font-size: 20px;
            font-weight: 800;
            """
        )

        content = QVBoxLayout()
        content.setSpacing(4)
        content.setContentsMargins(0, 2, 0, 2)

        self.title_label = QLabel(title)
        self.title_label.setStyleSheet(f"color: {MUTED}; font-size: 12px; font-weight: 600;")

        self.value_label = QLabel(value)
        self.value_label.setStyleSheet(f"color: {TEXT}; font-size: 18px; font-weight: 800;")

        content.addWidget(self.title_label)
        content.addWidget(self.value_label)
        content.addStretch()

        layout.addWidget(icon_box, 0, Qt.AlignmentFlag.AlignVCenter)
        layout.addLayout(content, 1)

        self.setStyleSheet(
            f"""
            QFrame#statCard {{
                background: {WHITE};
                border: 1px solid {BORDER};
                border-radius: 12px;
            }}
            """
        )

    def set_value(self, value: str):
        self.value_label.setText(str(value))


# ==========================================================
# دیالوگ جزئیات فاکتور
# ==========================================================
class ReportOrderDetailsDialog(QDialog):
    def __init__(self, order_data: dict, parent=None):
        super().__init__(parent)
        self.order_data = order_data

        self.setWindowTitle(f"جزئیات فاکتور {order_data['factor_number']}")
        self.resize(820, 600)
        self.setMinimumSize(700, 480)
        self.setLayoutDirection(Qt.LayoutDirection.RightToLeft)

        root = QVBoxLayout(self)
        root.setContentsMargins(22, 20, 22, 20)
        root.setSpacing(16)

        # هدر دیالوگ
        header_row = QHBoxLayout()
        title = QLabel(f"فاکتور فروش: {to_persian_digits(order_data['factor_number'])}")
        title.setStyleSheet(f"font-size: 18px; font-weight: 800; color: {TEXT};")
        header_row.addWidget(title)
        header_row.addStretch()

        is_paid = order_data["is_paid"]
        status_lbl = QLabel("پرداخت‌شده" if is_paid else "پرداخت‌نشده")
        status_lbl.setStyleSheet(
            f"""
            background: {GREEN_LIGHT if is_paid else DANGER_LIGHT};
            color: {GREEN if is_paid else DANGER};
            padding: 6px 14px;
            border-radius: 10px;
            font-size: 12px;
            font-weight: 700;
            """
        )
        header_row.addWidget(status_lbl)
        root.addLayout(header_row)

        # باکس اطلاعات کلی
        info_box = QFrame()
        info_box.setObjectName("infoBox")
        info_grid = QGridLayout(info_box)
        info_grid.setContentsMargins(16, 14, 16, 14)
        info_grid.setHorizontalSpacing(24)
        info_grid.setVerticalSpacing(10)

        meta = [
            ("تاریخ ثبت:", to_persian_digits(order_data["persian_date"])),
            ("ساعت ثبت:", to_persian_digits(order_data["persian_time"])),
            ("نام مشتری:", customer_text(order_data["customer_name"])),
            ("ثبت‌کننده:", order_data["performed_by"] or "نامشخص"),
        ]

        for i, (lbl, val) in enumerate(meta):
            r = i // 2
            c = (i % 2) * 2
            lbl_widget = QLabel(lbl)
            lbl_widget.setStyleSheet(f"color: {MUTED}; font-size: 12px; font-weight: 500;")
            val_widget = QLabel(val)
            val_widget.setStyleSheet(f"color: {TEXT}; font-size: 13px; font-weight: 700;")
            info_grid.addWidget(lbl_widget, r, c)
            info_grid.addWidget(val_widget, r, c + 1)

        root.addWidget(info_box)

        # جدول اقلام
        table_label = QLabel("اقلام سفارش داده شده")
        table_label.setStyleSheet(f"font-size: 13px; font-weight: 700; color: {TEXT};")
        root.addWidget(table_label)

        self.items_table = QTableWidget(0, 5)
        self.items_table.setHorizontalHeaderLabels(
            ["نام محصول", "قیمت واحد", "تعداد", "جمع ردیف", "دستور تهیه و مواد مصرفی"]
        )
        self.items_table.setLayoutDirection(Qt.LayoutDirection.RightToLeft)
        self.items_table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        self.items_table.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectRows)
        self.items_table.setSelectionMode(QAbstractItemView.SelectionMode.SingleSelection)
        self.items_table.setAlternatingRowColors(True)
        self.items_table.setShowGrid(False)
        self.items_table.verticalHeader().setVisible(False)
        self.items_table.verticalHeader().setDefaultSectionSize(40)

        header = self.items_table.horizontalHeader()
        header.setSectionResizeMode(0, QHeaderView.ResizeMode.Stretch)
        header.setSectionResizeMode(1, QHeaderView.ResizeMode.ResizeToContents)
        header.setSectionResizeMode(2, QHeaderView.ResizeMode.ResizeToContents)
        header.setSectionResizeMode(3, QHeaderView.ResizeMode.ResizeToContents)
        header.setSectionResizeMode(4, QHeaderView.ResizeMode.Stretch)

        for item in order_data["items"]:
            row = self.items_table.rowCount()
            self.items_table.insertRow(row)

            values = [
                item["menu_item_name"],
                money(item["unit_price"]),
                format_number(item["quantity"]),
                money(item["subtotal"]),
                item["recipe_snapshot_details"] or "—",
            ]

            for col, val in enumerate(values):
                cell = QTableWidgetItem(str(val))
                cell.setTextAlignment(
                    Qt.AlignmentFlag.AlignCenter if col in (1, 2, 3) else (Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter)
                )
                self.items_table.setItem(row, col, cell)

        root.addWidget(self.items_table, 1)

        # فوتر: جمع کل و دکمه بستن
        footer = QHBoxLayout()
        total_box = QHBoxLayout()
        total_box.setSpacing(10)

        items_count_lbl = QLabel(f"تعداد کل اقلام: <b>{format_number(order_data['total_items_count'])}</b>")
        items_count_lbl.setStyleSheet(f"color: {MUTED}; font-size: 13px;")

        total_price_lbl = QLabel(f"مبلغ قابل پرداخت: <b style='color: {GREEN};'>{money(order_data['total_price'])}</b>")
        total_price_lbl.setStyleSheet("font-size: 15px;")

        total_box.addWidget(items_count_lbl)
        total_box.addSpacing(16)
        total_box.addWidget(total_price_lbl)

        footer.addLayout(total_box)
        footer.addStretch()

        close_btn = QPushButton("بستن")
        close_btn.setMinimumSize(100, 36)
        close_btn.clicked.connect(self.accept)
        footer.addWidget(close_btn)

        root.addLayout(footer)

        self.setStyleSheet(
            f"""
            QDialog {{
                background: {BG};
                font-family: "Vazirmatn", "Noto Sans Arabic", "DejaVu Sans";
            }}
            QFrame#infoBox {{
                background: {WHITE};
                border: 1px solid {BORDER};
                border-radius: 10px;
            }}
            QTableWidget {{
                background: {WHITE};
                border: 1px solid {BORDER};
                border-radius: 10px;
                alternate-background-color: #F8FAF9;
                font-size: 12px;
            }}
            QHeaderView::section {{
                background: #EFF4F0;
                color: {MUTED};
                border: none;
                padding: 8px 6px;
                font-weight: 700;
                font-size: 11px;
            }}
            QTableWidget::item {{
                padding: 6px;
                border-bottom: 1px solid #EDF1EE;
            }}
            QPushButton {{
                background: {WHITE};
                color: {TEXT};
                border: 1px solid {BORDER};
                border-radius: 8px;
                font-weight: 600;
                padding: 6px 16px;
            }}
            QPushButton:hover {{
                background: #F0F4F2;
            }}
            """
        )


# ==========================================================
# صفحه گزارشات اصلی
# ==========================================================
class ReportsPage(QWidget):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.orders = []
        self.filtered_orders = []

        self._build_ui()
        self.refresh_reports()

    def _build_ui(self):
        self.setLayoutDirection(Qt.LayoutDirection.RightToLeft)

        root = QVBoxLayout(self)
        root.setContentsMargins(20, 18, 20, 20)
        root.setSpacing(16)

        # ─── ۱. هدر صفحه ──────────────────────────────────────────
        heading = QHBoxLayout()
        title_box = QVBoxLayout()
        title_box.setSpacing(4)

        title = QLabel("گزارش‌های فروش و فاکتورها")
        title.setStyleSheet(f"font-size: 20px; font-weight: 800; color: {TEXT};")

        subtitle = QLabel("مرور و بررسی سفارش‌های ثبت‌شده، جستجو و وضعیت تسویه‌حساب")
        subtitle.setStyleSheet(f"color: {MUTED}; font-size: 12px;")

        title_box.addWidget(title)
        title_box.addWidget(subtitle)
        heading.addLayout(title_box)
        heading.addStretch()

        # دکمه خروجی اکسل
        self.export_button = QPushButton("⬇ خروجی اکسل")
        self.export_button.setObjectName("exportButton")
        self.export_button.setCursor(Qt.CursorShape.PointingHandCursor)
        self.export_button.setMinimumHeight(38)
        self.export_button.setToolTip(
            "در صورت انتخاب یک یا چند فاکتور، همان‌ها خروجی گرفته می‌شوند؛\n"
            "در غیر این صورت از کل فاکتورهای فیلترشده خروجی تهیه می‌شود."
        )
        self.export_button.clicked.connect(self.export_to_excel)
        heading.addWidget(self.export_button)

        # دکمه به‌روزرسانی
        self.refresh_button = QPushButton("↻ به‌روزرسانی")
        self.refresh_button.setCursor(Qt.CursorShape.PointingHandCursor)
        self.refresh_button.setObjectName("primaryButton")
        self.refresh_button.setMinimumHeight(38)
        self.refresh_button.clicked.connect(self.refresh_reports)
        heading.addWidget(self.refresh_button)

        root.addLayout(heading)

        # ─── ۲. کارت‌های آمار کلیدی ────────────────────────────────
        stats_layout = QHBoxLayout()
        stats_layout.setSpacing(14)

        self.orders_stat = ReportStatCard("تعداد کل فاکتورها", "۰", "▤", "#2587C7")
        self.sales_stat = ReportStatCard("مجموع فروش", "۰ تومان", "▲", GREEN)
        self.paid_stat = ReportStatCard("تسویه‌شده", "۰ تومان", "✓", "#2E8B57")
        self.unpaid_stat = ReportStatCard("مطالبات باز (پرداخت‌نشده)", "۰ تومان", "!", DANGER)

        stats_layout.addWidget(self.orders_stat)
        stats_layout.addWidget(self.sales_stat)
        stats_layout.addWidget(self.paid_stat)
        stats_layout.addWidget(self.unpaid_stat)
        root.addLayout(stats_layout)

        # ─── ۳. پنل فیلتر و جستجو ─────────────────────────────────
        filter_panel = QFrame()
        filter_panel.setObjectName("filterPanel")
        filter_layout = QVBoxLayout(filter_panel)
        filter_layout.setContentsMargins(16, 14, 16, 14)
        filter_layout.setSpacing(12)

        # ردیف اول: جستجوی متنی و دراپ‌داون‌ها
        row_one = QHBoxLayout()
        row_one.setSpacing(10)

        self.search_input = QLineEdit()
        self.search_input.setPlaceholderText("جستجو در شماره فاکتور، نام مشتری یا ثبت‌کننده...")
        self.search_input.setClearButtonEnabled(True)
        self.search_input.setMinimumHeight(38)
        self.search_input.textChanged.connect(self.apply_filters)

        self.payment_filter = QComboBox()
        self.payment_filter.addItem("همه وضعیت‌های پرداخت", "all")
        self.payment_filter.addItem("فقط پرداخت‌شده", "paid")
        self.payment_filter.addItem("فقط پرداخت‌نشده", "unpaid")
        self.payment_filter.setMinimumHeight(38)
        self.payment_filter.currentIndexChanged.connect(self.apply_filters)

        self.customer_filter = QComboBox()
        self.customer_filter.addItem("همه مشتریان", "all")
        self.customer_filter.addItem("مشتریان ثبت‌شده", "registered")
        self.customer_filter.addItem("مشتریان مهمان", "guest")
        self.customer_filter.setMinimumHeight(38)
        self.customer_filter.currentIndexChanged.connect(self.apply_filters)

        row_one.addWidget(self.search_input, 4)
        row_one.addWidget(self.payment_filter, 2)
        row_one.addWidget(self.customer_filter, 2)
        filter_layout.addLayout(row_one)

        # ردیف دوم: بازه زمانی + دکمه‌ها
        row_two = QHBoxLayout()
        row_two.setSpacing(10)

        date_lbl = QLabel("بازه تاریخ شمسی:")
        date_lbl.setStyleSheet(f"color: {MUTED}; font-size: 12px; font-weight: 600;")
        row_two.addWidget(date_lbl)

        self.date_from_input = QLineEdit()
        self.date_from_input.setPlaceholderText("از: 1403/01/01")
        self.date_from_input.setMinimumHeight(36)
        self.date_from_input.setFixedWidth(140)

        self.date_to_input = QLineEdit()
        self.date_to_input.setPlaceholderText("تا: 1403/12/29")
        self.date_to_input.setMinimumHeight(36)
        self.date_to_input.setFixedWidth(140)

        row_two.addWidget(self.date_from_input)
        row_two.addWidget(self.date_to_input)

        self.apply_date_button = QPushButton("اعمال تاریخ")
        self.apply_date_button.setMinimumHeight(36)
        self.apply_date_button.clicked.connect(self.apply_filters)
        row_two.addWidget(self.apply_date_button)

        self.clear_filters_button = QPushButton("پاک‌سازی فیلترها")
        self.clear_filters_button.setMinimumHeight(36)
        self.clear_filters_button.clicked.connect(self.clear_filters)
        row_two.addWidget(self.clear_filters_button)

        row_two.addStretch()

        self.result_count = QLabel("۰ فاکتور یافت شد")
        self.result_count.setStyleSheet(f"color: {MUTED}; font-size: 12px; font-weight: 700;")
        row_two.addWidget(self.result_count)

        filter_layout.addLayout(row_two)
        root.addWidget(filter_panel)

        # ─── ۴. جدول نتایج ────────────────────────────────────────
        table_panel = QFrame()
        table_panel.setObjectName("tablePanel")
        table_layout = QVBoxLayout(table_panel)
        table_layout.setContentsMargins(14, 14, 14, 14)
        table_layout.setSpacing(10)

        self.reports_table = QTableWidget(0, 7)
        self.reports_table.setHorizontalHeaderLabels(
            ["شماره فاکتور", "تاریخ", "زمان", "مشتری", "تعداد اقلام", "مبلغ کل", "وضعیت پرداخت"]
        )
        self.reports_table.setLayoutDirection(Qt.LayoutDirection.RightToLeft)
        self.reports_table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        self.reports_table.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectRows)
        # پشتیبانی از انتخاب چند فاکتور هم‌زمان با Ctrl و Shift
        self.reports_table.setSelectionMode(QAbstractItemView.SelectionMode.ExtendedSelection)
        self.reports_table.setAlternatingRowColors(True)
        self.reports_table.setShowGrid(False)
        self.reports_table.verticalHeader().setVisible(False)
        self.reports_table.verticalHeader().setDefaultSectionSize(44)

        header = self.reports_table.horizontalHeader()
        header.setSectionResizeMode(0, QHeaderView.ResizeMode.Stretch)
        header.setSectionResizeMode(1, QHeaderView.ResizeMode.ResizeToContents)
        header.setSectionResizeMode(2, QHeaderView.ResizeMode.ResizeToContents)
        header.setSectionResizeMode(3, QHeaderView.ResizeMode.Stretch)
        header.setSectionResizeMode(4, QHeaderView.ResizeMode.ResizeToContents)
        header.setSectionResizeMode(5, QHeaderView.ResizeMode.ResizeToContents)
        header.setSectionResizeMode(6, QHeaderView.ResizeMode.ResizeToContents)

        self.reports_table.cellDoubleClicked.connect(self.open_selected_order)
        table_layout.addWidget(self.reports_table, 1)

        # فوتر جدول: راهنما و دکمه جزئیات
        table_footer = QHBoxLayout()
        hint = QLabel("💡 با نگه داشتن کلید Ctrl یا Shift می‌توانید چند فاکتور را انتخاب کرده و خروجی بگیرید. برای مشاهده ریزاقلام، روی ردیف دوبار کلیک کنید.")
        hint.setStyleSheet(f"color: {MUTED}; font-size: 11px;")
        table_footer.addWidget(hint)
        table_footer.addStretch()

        self.details_button = QPushButton("مشاهده فاکتور انتخاب‌شده")
        self.details_button.setObjectName("primaryButton")
        self.details_button.setMinimumHeight(38)
        self.details_button.clicked.connect(self.open_selected_order)
        table_footer.addWidget(self.details_button)

        table_layout.addLayout(table_footer)
        root.addWidget(table_panel, 1)

        # ─── ۵. شیوه ظاهری جامع (QSS) ─────────────────────────────
        self.setStyleSheet(
            f"""
            QWidget {{
                color: {TEXT};
                font-family: "Vazirmatn", "Noto Sans Arabic", "DejaVu Sans";
                font-size: 12px;
            }}
            QFrame#filterPanel, QFrame#tablePanel {{
                background: {WHITE};
                border: 1px solid {BORDER};
                border-radius: 12px;
            }}
            QLineEdit, QComboBox {{
                background: {WHITE};
                border: 1px solid {BORDER};
                border-radius: 8px;
                padding: 6px 12px;
                color: {TEXT};
                selection-background-color: {GREEN_LIGHT};
                selection-color: {TEXT};
            }}
            QLineEdit:focus, QComboBox:focus {{
                border: 1.5px solid {GREEN};
            }}
            QComboBox::drop-down {{
                border: none;
                width: 24px;
            }}
            QPushButton {{
                background: {WHITE};
                color: {TEXT};
                border: 1px solid {BORDER};
                border-radius: 8px;
                padding: 6px 14px;
                font-weight: 600;
            }}
            QPushButton:hover {{
                background: #F2F7F4;
                border-color: #D3DFD7;
            }}
            QPushButton#primaryButton {{
                background: {GREEN};
                color: {WHITE};
                border: none;
                font-weight: 700;
                padding: 6px 18px;
            }}
            QPushButton#primaryButton:hover {{
                background: {GREEN_DARK};
            }}
            QPushButton#exportButton {{
                background: {GREEN_LIGHT};
                color: {GREEN};
                border: 1px solid #BFD9C7;
                border-radius: 8px;
                padding: 6px 16px;
                font-weight: 700;
            }}
            QPushButton#exportButton:hover {{
                background: #DCEEE2;
                border-color: #93BDA2;
            }}
            QPushButton#exportButton:disabled {{
                background: #F1F5F2;
                color: #94A39A;
                border-color: #E2E8E4;
            }}
            QTableWidget {{
                background: {WHITE};
                border: none;
                alternate-background-color: #F8FAF9;
                selection-background-color: {GREEN_LIGHT};
                selection-color: {TEXT};
            }}
            QHeaderView::section {{
                background: #F1F6F3;
                color: {MUTED};
                border: none;
                padding: 10px 8px;
                font-weight: 700;
                font-size: 11px;
            }}
            QTableWidget::item {{
                padding: 6px 10px;
                border-bottom: 1px solid #EDF2EF;
            }}
            """
        )

    # ------------------------------------------------------
    # خواندن اطلاعات از دیتابیس
    # ------------------------------------------------------
    def refresh_reports(self):
        self.refresh_button.setEnabled(False)
        session = init_db()

        try:
            db_orders = (
                session.query(OrderArchive)
                .order_by(OrderArchive.id.desc())
                .all()
            )

            self.orders = [
                {
                    "id": order.id,
                    "factor_number": order.factor_number or "",
                    "total_price": float(order.total_price or 0),
                    "performed_by": order.performed_by or "",
                    "customer_id": order.customer_id,
                    "customer_name": order.customer_name,
                    "total_items_count": int(order.total_items_count or 0),
                    "is_paid": bool(order.is_paid),
                    "persian_date": order.persian_date or "",
                    "persian_time": order.persian_time or "",
                }
                for order in db_orders
            ]
        except SQLAlchemyError as error:
            self.orders = []
            QMessageBox.critical(
                self,
                "خطا در بارگذاری گزارشات",
                f"خطایی هنگام خواندن داده‌ها از دیتابیس رخ داد:\n{error}",
            )
        finally:
            session.close()
            self.refresh_button.setEnabled(True)

        self.apply_filters()

    # ------------------------------------------------------
    # فیلتر و جستجو
    # ------------------------------------------------------
    def apply_filters(self):
        search_text = self.search_input.text().strip().casefold()
        payment_status = self.payment_filter.currentData()
        customer_status = self.customer_filter.currentData()

        date_from = self.date_from_input.text().strip()
        date_to = self.date_to_input.text().strip()

        if date_from and not self._valid_date(date_from):
            self.result_count.setText("فرمت تاریخ شروع نامعتبر است (مثال: 1403/01/01)")
            return

        if date_to and not self._valid_date(date_to):
            self.result_count.setText("فرمت تاریخ پایان نامعتبر است (مثال: 1403/12/29)")
            return

        if date_from and date_to and date_from > date_to:
            self.result_count.setText("تاریخ شروع نباید بعد از تاریخ پایان باشد")
            return

        filtered = []

        for order in self.orders:
            searchable = " ".join([
                order["factor_number"],
                order["customer_name"] or "",
                order["performed_by"],
            ]).casefold()

            if search_text and search_text not in searchable:
                continue

            if payment_status == "paid" and not order["is_paid"]:
                continue
            if payment_status == "unpaid" and order["is_paid"]:
                continue

            if customer_status == "registered" and order["customer_id"] is None:
                continue
            if customer_status == "guest" and order["customer_id"] is not None:
                continue

            order_date = order["persian_date"]
            if date_from and order_date < date_from:
                continue
            if date_to and order_date > date_to:
                continue

            filtered.append(order)

        self.filtered_orders = filtered
        self.populate_table(filtered)
        self.update_statistics(filtered)

    @staticmethod
    def _valid_date(value: str) -> bool:
        parts = value.split("/")
        if len(parts) != 3 or not all(p.isdigit() for p in parts):
            return False
        year, month, day = map(int, parts)
        return year >= 1300 and 1 <= month <= 12 and 1 <= day <= 31

    def clear_filters(self):
        self.search_input.clear()
        self.payment_filter.setCurrentIndex(0)
        self.customer_filter.setCurrentIndex(0)
        self.date_from_input.clear()
        self.date_to_input.clear()
        self.apply_filters()

    # ------------------------------------------------------
    # پر کردن جدول
    # ------------------------------------------------------
    def populate_table(self, orders: list[dict]):
        self.reports_table.setRowCount(0)

        for row, order in enumerate(orders):
            self.reports_table.insertRow(row)

            cells = [
                (order["factor_number"], Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter),
                (to_persian_digits(order["persian_date"]), Qt.AlignmentFlag.AlignCenter),
                (to_persian_digits(order["persian_time"]), Qt.AlignmentFlag.AlignCenter),
                (customer_text(order["customer_name"]), Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter),
                (format_number(order["total_items_count"]), Qt.AlignmentFlag.AlignCenter),
                (money(order["total_price"]), Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter),
            ]

            for col, (val, align) in enumerate(cells):
                cell_item = QTableWidgetItem(str(val))
                cell_item.setTextAlignment(align)
                if col == 0:
                    cell_item.setData(Qt.ItemDataRole.UserRole, order["id"])
                self.reports_table.setItem(row, col, cell_item)

            # ستون وضعیت پرداخت با Badge اختصاصی
            badge = PaymentBadge(order["is_paid"])
            self.reports_table.setCellWidget(row, 6, badge)

        self.result_count.setText(f"{format_number(len(orders))} فاکتور یافت شد")

    # ------------------------------------------------------
    # خروجی اکسل (پشتیبانی از تک، چند انتخابی و کل نتایج فیلتر شده)
    # ------------------------------------------------------
    def export_to_excel(self):
        # بررسی ردیف‌های انتخاب‌شده
        selected_rows = sorted({idx.row() for idx in self.reports_table.selectedIndexes()})

        if selected_rows:
            selected_ids = []
            for row in selected_rows:
                item = self.reports_table.item(row, 0)
                if item:
                    oid = item.data(Qt.ItemDataRole.UserRole)
                    if oid is not None:
                        selected_ids.append(oid)
            
            selected_id_set = set(selected_ids)
            orders_to_export = [o for o in self.filtered_orders if o["id"] in selected_id_set]
        else:
            orders_to_export = list(self.filtered_orders)

        if not orders_to_export:
            QMessageBox.information(
                self,
                "عدم وجود داده",
                "هیچ فاکتوری برای دریافت خروجی اکسل در لیست فیلترشده موجود نیست."
            )
            return

        file_path, _ = QFileDialog.getSaveFileName(
            self,
            "ذخیره خروجی اکسل فاکتورها",
            "گزارش_فاکتورها.xlsx",
            "Excel Files (*.xlsx)"
        )

        if not file_path:
            return

        if not file_path.lower().endswith(".xlsx"):
            file_path += ".xlsx"

        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment
            from openpyxl.utils import get_column_letter
        except ImportError:
            QMessageBox.critical(
                self,
                "نیاز به نصب کتابخانه",
                "برای تهیه خروجی اکسل، کتابخانه openpyxl لازم است.\nلطفاً دستور زیر را در ترمینال اجرا کنید:\npip install openpyxl"
            )
            return

        session = None
        try:
            session = init_db()
            target_ids = [o["id"] for o in orders_to_export]

            db_orders = (
                session.query(OrderArchive)
                .filter(OrderArchive.id.in_(target_ids))
                .all()
            )

            # مرتب‌سازی مطابق با ترتیب جدول
            order_map = {order.id: order for order in db_orders}
            sorted_orders = [order_map[oid] for oid in target_ids if oid in order_map]

            if not sorted_orders:
                QMessageBox.warning(self, "خطا", "اطلاعات فاکتورهای انتخاب‌شده یافت نشد.")
                return

            wb = Workbook()

            # شیت اول: لیست فاکتورها
            ws_orders = wb.active
            ws_orders.title = "فاکتورها"
            ws_orders.append([
                "شماره فاکتور",
                "تاریخ ثبت",
                "ساعت ثبت",
                "نام مشتری",
                "ثبت‌کننده",
                "تعداد اقلام",
                "مبلغ کل (تومان)",
                "وضعیت پرداخت"
            ])

            # شیت دوم: ریز اقلام
            ws_items = wb.create_sheet(title="اقلام فاکتورها")
            ws_items.append([
                "شماره فاکتور",
                "تاریخ ثبت",
                "نام مشتری",
                "نام محصول",
                "تعداد",
                "قیمت واحد (تومان)",
                "جمع ردیف (تومان)",
                "دستور تهیه و مواد مصرفی"
            ])

            for order in sorted_orders:
                f_num = order.factor_number or ""
                p_date = order.persian_date or ""
                p_time = order.persian_time or ""
                c_name = customer_text(order.customer_name)
                p_by = order.performed_by or "نامشخص"

                ws_orders.append([
                    f_num,
                    p_date,
                    p_time,
                    c_name,
                    p_by,
                    int(order.total_items_count or 0),
                    float(order.total_price or 0),
                    "پرداخت‌شده" if order.is_paid else "پرداخت‌نشده"
                ])

                for item in order.items:
                    ws_items.append([
                        f_num,
                        p_date,
                        c_name,
                        item.menu_item_name or "",
                        int(item.quantity or 0),
                        float(item.unit_price or 0),
                        float(item.subtotal or 0),
                        item.recipe_snapshot_details or "—"
                    ])

            # استایل‌دهی و قالب‌بندی شیت‌ها
            def style_sheet(ws, col_widths, currency_cols=()):
                ws.sheet_view.rightToLeft = True
                ws.freeze_panes = "A2"
                ws.auto_filter.ref = ws.dimensions
                ws.row_dimensions[1].height = 28

                header_fill = PatternFill(start_color="28745A", end_color="28745A", fill_type="solid")
                header_font = Font(name="Vazirmatn", size=11, bold=True, color="FFFFFF")
                body_font = Font(name="Vazirmatn", size=10)

                for cell in ws[1]:
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

                for row in ws.iter_rows(min_row=2):
                    for cell in row:
                        cell.font = body_font
                        cell.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)

                    for col_idx in currency_cols:
                        row[col_idx - 1].number_format = '#,##0'

                for idx, w in enumerate(col_widths, start=1):
                    ws.column_dimensions[get_column_letter(idx)].width = w

            style_sheet(
                ws_orders,
                col_widths=[20, 16, 12, 26, 20, 14, 22, 18],
                currency_cols=(7,)
            )

            style_sheet(
                ws_items,
                col_widths=[20, 16, 26, 28, 12, 20, 20, 55],
                currency_cols=(6, 7)
            )

            wb.save(file_path)

            QMessageBox.information(
                self,
                "خروجی با موفقیت ذخیره شد",
                f"فایل اکسل شامل {format_number(len(sorted_orders))} فاکتور با موفقیت ذخیره شد:\n{file_path}"
            )

        except SQLAlchemyError as error:
            QMessageBox.critical(self, "خطا در دیتابیس", f"خطایی در خواندن جزئیات فاکتورها رخ داد:\n{error}")
        except Exception as error:
            QMessageBox.critical(self, "خطا در ساخت اکسل", f"ساخت خروجی با خطا مواجه شد:\n{error}")
        finally:
            if session:
                session.close()

    # ------------------------------------------------------
    # آمار کارت‌ها
    # ------------------------------------------------------
    def update_statistics(self, orders: list[dict]):
        total_count = len(orders)
        total_sales = sum(o["total_price"] for o in orders)
        paid_sales = sum(o["total_price"] for o in orders if o["is_paid"])
        unpaid_sales = sum(o["total_price"] for o in orders if not o["is_paid"])

        self.orders_stat.set_value(format_number(total_count))
        self.sales_stat.set_value(money(total_sales))
        self.paid_stat.set_value(money(paid_sales))
        self.unpaid_stat.set_value(money(unpaid_sales))

    # ------------------------------------------------------
    # باز کردن جزئیات
    # ------------------------------------------------------
    def open_selected_order(self, row: int = None, column: int = None):
        if row is None or isinstance(row, bool):
            row = self.reports_table.currentRow()

        if row < 0:
            QMessageBox.information(self, "انتخاب فاکتور", "لطفاً ابتدا یک فاکتور را از جدول انتخاب نمایید.")
            return

        item = self.reports_table.item(row, 0)
        if not item:
            return

        order_id = item.data(Qt.ItemDataRole.UserRole)
        if not order_id:
            return

        session = init_db()
        try:
            order = session.query(OrderArchive).filter(OrderArchive.id == order_id).first()
            if not order:
                QMessageBox.warning(self, "یافت نشد", "فاکتور موردنظر پیدا نشد.")
                return

            order_data = {
                "factor_number": order.factor_number,
                "total_price": float(order.total_price or 0),
                "performed_by": order.performed_by or "",
                "customer_id": order.customer_id,
                "customer_name": order.customer_name,
                "total_items_count": int(order.total_items_count or 0),
                "is_paid": bool(order.is_paid),
                "persian_date": order.persian_date or "",
                "persian_time": order.persian_time or "",
                "items": [
                    {
                        "menu_item_name": itm.menu_item_name,
                        "unit_price": float(itm.unit_price or 0),
                        "quantity": int(itm.quantity or 0),
                        "subtotal": float(itm.subtotal or 0),
                        "recipe_snapshot_details": itm.recipe_snapshot_details or "",
                    }
                    for itm in order.items
                ],
            }
        except SQLAlchemyError as error:
            QMessageBox.critical(self, "خطا", f"خطا در خواندن جزئیات فاکتور:\n{error}")
            return
        finally:
            session.close()

        dialog = ReportOrderDetailsDialog(order_data, self)
        dialog.exec()
